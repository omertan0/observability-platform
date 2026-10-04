from pathlib import Path
from docx import Document
from docx.shared import RGBColor

root = Path(__file__).resolve().parent
path = root / 'rapor.docx'
doc = Document(path)
doc.paragraphs[1].text = 'Araştırma Mimari Tasarım ve İlerleme Raporu'
anchor = next(p for p in doc.paragraphs if p.text.strip() == 'Kaynakça')

def add(text, style='Normal', page_break=False):
    p = anchor.insert_paragraph_before(text, style=style)
    p.paragraph_format.page_break_before = page_break
    if style.startswith('Heading'):
        for run in p.runs:
            run.font.color.rgb = RGBColor(0, 0, 0)
    return p

add('9 Proje ilerleme durumu', 'Heading 1', True)
add('5 Ekim 2026')
add('Projenin Faz 1 araştırma ve tasarım çalışmaları hazırlanmış, Faz 2 kapsamında veri alım API’sinin ilk sürümü geliştirilmiştir. API, log ve metrik kayıtlarını doğrulayıp alınma zamanını ekleyerek yanıt üretmektedir. Redis servisi Docker üzerinden çalıştırılmıştır. Celery bağlantı ayarı ve ilk görev tanımı hazırlanmış; worker’ın çalıştırılması ve API ile kuyruk arasındaki bağlantı bir sonraki uygulama adımı olarak belirlenmiştir.')
add('9.1 Tamamlanan çalışmalar', 'Heading 2')
add('Logs, Metrics ve Traces kavramları ile OpenTelemetry, Elastic APM ve Prometheus mimarileri araştırılmıştır. Python ve FastAPI, Redis ve Celery, PostgreSQL ve TimescaleDB, React ve SSE bileşenlerinden oluşan mimari seçilmiştir. Log ve metrik tabloları, günlük zaman parçaları ve sorgu indeksleri SQL dosyasında tanımlanmıştır. JSON kayıt örnekleri ve veri alım API sözleşmesi hazırlanmıştır.')
add('Python sanal ortamı oluşturulmuş ve geliştirme paketleri kurulmuştur. FastAPI ile GET /health, POST /logs ve POST /metrics uç noktaları yazılmıştır. Pydantic modelleri zorunlu alanları, log seviyelerini, boş metinleri, saat dilimi içeren zamanları ve sonlu sayısal metrik değerlerini kontrol etmektedir. API, received_at alanını UTC olarak eklemektedir. OpenAPI tanımı ve Swagger deneme ekranı erişilebilir durumdadır.')
add('WSL ve Docker Desktop kurulmuştur. Redis servisi compose.yaml ile tanımlanmış; yerel bağlantı noktası, dosyaya kayıt ayarı ve kalıcı depolama alanı yapılandırılmıştır. Celery’nin Redis bağlantı ayarı ile alınan logu terminalde gösteren process_log görevi yazılmıştır.')
add('9.2 Yapılan kontroller', 'Heading 2')
add('Manuel API denemelerinde geçerli log ve metrik kayıtları için 200 yanıtı, geçersiz log seviyesi ve metin olarak gönderilen metrik değeri için 422 yanıtı görülmüştür. Docker sürüm kontrolünde hem istemci hem sunucu bilgileri alınmıştır. Redis konteynerinde PING komutuna PONG yanıtı alınarak servisin cevap verdiği doğrulanmıştır. Yük testi ve uçtan uca kuyruk işleme testi henüz yapılmamıştır.')
add('9.3 Devam eden ve planlanan çalışmalar', 'Heading 2')
add('Mevcut API kayıtları doğrulayıp yanıtta göstermektedir; henüz kuyruğa gönderme veya veritabanına yazma işlemi yapmamaktadır. Sıradaki çalışmalar Celery worker’ın Docker içinde çalıştırılması, API’den görev gönderilmesi ve PostgreSQL ile TimescaleDB kurulumunun tamamlanmasıdır. SQL şeması daha sonra veritabanına uygulanacaktır. Kuyruğa alınan kayıtlar için sözleşmede planlanan 202 yanıtı bu bağlantı kurulduğunda kullanılacaktır.')
add('Veri kaydetme akışı tamamlandıktan sonra sorgulama uç noktaları ve alarm motoru geliştirilecektir. Faz 3 kapsamında React paneli ve SSE ile canlı akış; Faz 4 kapsamında log simülatörü, yük testleri, sistemin Docker Compose ile bütünleştirilmesi, canlıya alınması ve teslim belgeleri hazırlanacaktır.')

doc.save(path)
print('İlerleme bölümü rapora eklendi.')
